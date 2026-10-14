"""Slides for one lesson. Text is merged with the picture or code it explains,
so one slide = one idea. Animations play as GIFs."""

import json
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from lxml import etree
from PIL import Image, ImageChops
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from pages import homework_tasks

ANIMS_DIR = Path(__file__).resolve().parent / "anims"
GIF_SCRIPT = Path(__file__).resolve().parent / "gif.mjs"

NAVY = RGBColor(0x1F, 0x2A, 0x37)
BLUE = RGBColor(0x1D, 0x5F, 0xD1)
GREY = RGBColor(0x5B, 0x66, 0x76)
CODE_BG = RGBColor(0x1E, 0x25, 0x30)
CODE_INK = RGBColor(0xE9, 0xED, 0xF2)

SLIDE_H = Inches(7.5)
LEFT = Inches(0.6)
WIDTH = Inches(12.13)
FADE = '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med"><p:fade/></p:transition>'


def plain(text):
    return text.replace("`", "").replace("**", "")


# ------------------------------------------------------------ animations --

def anim_seconds(name):
    svg = (ANIMS_DIR / f"{name}.html").read_text(encoding="utf-8")
    found = re.search(r'data-seconds="([\d.]+)"', svg)
    assert found, f"anims/{name}.html needs data-seconds on the <svg>"
    return float(found.group(1))


def anim_gif(name):
    """Animated GIF of an animation, made once and redrawn when the svg changes."""
    src, gif = ANIMS_DIR / f"{name}.html", ANIMS_DIR / f"{name}.gif"
    if gif.exists() and gif.stat().st_mtime >= src.stat().st_mtime:
        return gif
    if not shutil.which("node"):
        print("  no node, slide animation skipped:", name)
        return None
    with tempfile.TemporaryDirectory() as tmp:
        page = Path(tmp) / "a.html"
        page.write_text('<!DOCTYPE html><meta charset="UTF-8"><style>body{margin:0;background:#fff}'
                        'svg{display:block;width:720px}</style>'
                        f'<div class="anim">{src.read_text(encoding="utf-8")}</div>', encoding="utf-8")
        subprocess.run(["node", str(GIF_SCRIPT), str(page), tmp, str(anim_seconds(name))],
                       check=True, capture_output=True)
        times = json.loads((Path(tmp) / "times.json").read_text())
        frames, durations = [], []
        for i, f in enumerate(sorted(Path(tmp).glob("f*.png"))):
            img = Image.open(f).convert("RGB")
            ms = (times[i + 1] - times[i]) if i + 1 < len(times) else 2500
            if frames and not ImageChops.difference(img, frames[-1]).getbbox():
                durations[-1] += ms  # same picture: hold the last frame longer
                continue
            frames.append(img)
            durations.append(ms)
        durations[-1] = max(durations[-1], 2500)
        frames = [f.quantize(colors=128) for f in frames]
        frames[0].save(gif, save_all=True, append_images=frames[1:], duration=durations, loop=0, optimize=True)
    return gif


# ------------------------------------------------------------------ deck --

class Deck:
    def __init__(self, lesson_n):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.33), SLIDE_H
        self.lesson_n = lesson_n

    def band(self, s):
        """Blue strip on top and the lesson number at the bottom of every slide."""
        bar = s.shapes.add_shape(1, 0, 0, self.prs.slide_width, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = BLUE
        bar.line.fill.background()
        foot = s.shapes.add_textbox(self.prs.slide_width - Inches(2.2), SLIDE_H - Inches(0.45), Inches(1.9), Inches(0.35))
        p = foot.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = f"Lesson {self.lesson_n}"
        r.font.size, r.font.color.rgb = Pt(12), GREY

    def slide(self, title=None):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        s._element.append(etree.fromstring(FADE))
        self.band(s)
        if title:
            self.text(s, [(plain(title), 36, BLUE, True)], Inches(0.35), Inches(0.9))
        return s

    def text(self, s, lines, top, height, fill=None, mono=False):
        """lines: (text, size, color, bold)."""
        if fill:
            box = s.shapes.add_shape(1, LEFT, top, WIDTH, height)
            box.fill.solid()
            box.fill.fore_color.rgb = fill
            box.line.fill.background()
        else:
            box = s.shapes.add_textbox(LEFT, top, WIDTH, height)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        if fill:
            tf.margin_left = tf.margin_right = Inches(0.25)
            tf.margin_top = Inches(0.15)
        for i, (line, size, color, bold) in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.space_after = Pt(2 if mono else 10)
            r = p.add_run()
            r.text = line or " "
            r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color
            if mono:
                r.font.name = "Menlo"
        return box

    def save(self, path):
        self.prs.save(path)


TEXT_STYLE = {
    "say": (26, NAVY, False),
    "big": (30, NAVY, True),
    "uz": (20, GREY, False),
}


def text_lines(blocks):
    lines = []
    for kind, value in blocks:
        size, color, bold = TEXT_STYLE[kind]
        lines.append((("UZ: " if kind == "uz" else "") + plain(value), size, color, bold))
    return lines


def text_height(lines):
    """Rough height: about 70 characters fit on one line at 26pt."""
    inches = sum((len(t) // max(1, int(70 * 26 / size)) + 1) * size * 1.45 / 72 for t, size, _, _ in lines)
    return Inches(inches + 0.2)


def put_visual(d, s, b, top):
    kind, args = b[0], b[1:]
    room = SLIDE_H - top - Inches(0.3)
    if kind == "anim":
        media = anim_gif(args[0])
        if media:
            w = min(WIDTH, Emu(int(room * 720 / 260)))
            s.shapes.add_picture(str(media), LEFT + (WIDTH - w) // 2, top, width=w)
        d.text(s, [(plain(args[1]), 18, GREY, False)], SLIDE_H - Inches(0.6), Inches(0.5))
    elif kind == "code":
        lines = args[0].split("\n")
        box = room - Inches(0.55)
        size = max(10, min(22, int(box / 12700 / (len(lines) * 1.5))))  # 12700 EMU = 1 pt
        d.text(s, [(plain(args[1]), 20, GREY, True)], top, Inches(0.5))
        d.text(s, [(line, size, CODE_INK, False) for line in lines], top + Inches(0.55),
               min(box, Emu(int(12700 * size * 1.3 * len(lines))) + Inches(0.4)), fill=CODE_BG, mono=True)
    elif kind == "table":
        head, rows = args
        row_h = min(Inches(0.75), Emu(room // (len(rows) + 1)))
        shape = s.shapes.add_table(len(rows) + 1, len(head), LEFT, top, WIDTH, row_h * (len(rows) + 1))
        for r, row in enumerate([head] + rows):
            for c, cell in enumerate(row):
                run = shape.table.cell(r, c).text_frame.paragraphs[0].add_run()
                run.text = plain(cell)
                run.font.size = Pt(22)
    elif kind == "quiz":
        options = [(f"{'ABCD'[i]})  {plain(o)}", 26, NAVY, False) for i, o in enumerate(args[1])]
        d.text(s, [("Quick question: " + plain(args[0]), 28, NAVY, True)] + options, top, room)
    elif kind == "steps":
        d.text(s, [(f"{i}.  {plain(t)}", 24, NAVY, False) for i, t in enumerate(args[0], 1)], top, room)


def numbered(items, size):
    return [(f"{i}.  {plain(t)}", size, NAVY, False) for i, t in enumerate(items, 1)]


def build_slides(L, path):
    d = Deck(L["n"])
    s = d.slide()
    d.text(s, [(f"Lesson {L['n']}", 30, GREY, True), (plain(L["title"]), 54, BLUE, True),
               (plain(L["goal"]), 26, NAVY, False)], Inches(2.2), Inches(3.5))
    if L["review"]:
        lines = []
        for i, (_, q, options, _) in enumerate(L["review"], 1):
            lines.append((f"{i}. {plain(q)}", 24, NAVY, True))
            lines.append(("     " + "     ".join(f"{'ABCD'[k]}) {plain(o)}" for k, o in enumerate(options)), 20, GREY, False))
        d.text(d.slide("Warm-up"), lines, Inches(1.4), Inches(5.8))

    part, pending = None, []

    def flush():
        if pending:
            d.text(d.slide(part), text_lines(pending), Inches(1.4), Inches(5.6))
            pending.clear()

    for b in L["blocks"]:
        kind = b[0]
        if kind == "h":
            flush()
            part = b[1]
        elif kind in TEXT_STYLE:
            pending.append(b)
        elif kind in ("try", "teacher"):
            continue  # done live on the lesson page
        else:
            lines = text_lines(pending)
            long_code = kind == "code" and b[1].count("\n") > 12
            if lines and (long_code or text_height(lines) > Inches(1.8)):
                flush()
                lines = []
            s = d.slide(part)
            top = Inches(1.35)
            if lines:
                d.text(s, lines, top, text_height(lines))
                top += text_height(lines)
            pending.clear()
            put_visual(d, s, b, top)
    flush()
    d.text(d.slide("Watch out!"), [(plain(m), 24, NAVY, False) for m in L["mistakes"]], Inches(1.4), Inches(5.6))
    s = d.slide("Practice and homework")
    d.text(s, [("Practice in class", 24, BLUE, True)] + numbered([t["task"] for t in L["practice"]], 18)
           + [("Homework", 24, BLUE, True)] + numbered([t["task"] for t in homework_tasks(L)], 18),
           Inches(1.3), Inches(6))
    d.save(path)
